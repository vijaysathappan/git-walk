import io
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from app import database
from app.euc.dependency.services import DependencyService
from app.euc.intelligence.services import IntelligenceService
from app.euc.migration.services import MigrationService
from app.euc.service import analyze_euc, ingest_euc
from app.services.semantic_ledger_service import ledger_for_connection


def migration_workbook() -> bytes:
    workbook = Workbook()
    inputs = workbook.active
    inputs.title = "Master Data"
    inputs.append(["Product", "Price"])
    inputs.append(["A", 10])
    inputs.append(["B", 20])
    table = Table(displayName="ProductTable", ref="A1:B3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    inputs.add_table(table)
    pricing = workbook.create_sheet("Pricing")
    pricing.append(["Quantity", "Product Price", "Revenue"])
    pricing.append([3, "='Master Data'!B2", "=A2*B2"])
    report = workbook.create_sheet("Report")
    report.append(["Revenue", "Dynamic"])
    report.append(["=Pricing!C2", '=INDIRECT("A"&2)'])
    output = io.BytesIO(); workbook.save(output); workbook.close()
    return output.getvalue()


class MigrationServiceIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_db_path = database.DB_PATH
        database.DB_PATH = Path(self.temp_dir.name) / "migration.db"
        database.initialize_product_schema()
        self.user = database.get_or_create_user("migration-owner@example.com")
        database.create_sqlite_table_from_df("QUEUE_BOARD_MIG", pd.DataFrame([{"VALUE": 1}]))
        registered = database.register_dataset("QUEUE_BOARD_MIG", self.user["user_id"], "source.xlsx", 1, 1)
        self.repository_id = registered["repository_id"]
        self.asset = ingest_euc(self.repository_id, "migration.xlsx", migration_workbook(), self.user["user_id"])
        analyze_euc(self.asset["euc_id"], self.user["user_id"])
        DependencyService().build(self.asset["euc_id"], self.user["user_id"])
        IntelligenceService().build(self.asset["euc_id"], self.user["user_id"])

    def tearDown(self):
        database.DB_PATH = self.original_db_path
        self.temp_dir.cleanup()

    def test_blueprint_is_deterministic_traceable_and_queryable(self):
        service = MigrationService()
        first = service.analyze(self.asset["euc_id"], self.user["user_id"])
        second = service.analyze(self.asset["euc_id"], self.user["user_id"])
        self.assertEqual(first["blueprint_manifest_hash"], second["blueprint_manifest_hash"])
        self.assertEqual("CURRENT", second["freshness"])
        self.assertGreater(second["unit_count"], 0)
        self.assertGreaterEqual(second["readiness"]["score"], 0)
        self.assertLessEqual(second["readiness"]["score"], 100)
        self.assertAlmostEqual(100, sum(second["coverage"].values()), places=1)
        self.assertIn(second["strategy"], {"LIFT_AND_GOVERN", "HYBRID_MODERNIZATION",
                                           "NATIVE_REBUILD", "DECOMPOSE_AND_MIGRATE",
                                           "RETAIN_IN_EXCEL", "RETIRE"})
        components = service.components(self.asset["euc_id"], self.user["user_id"])
        self.assertTrue(any(item["source_type"] == "FORMULA_PATTERN" for item in components["items"]))
        self.assertTrue(service.waves(self.asset["euc_id"], self.user["user_id"])["items"])
        self.assertTrue(service.target(self.asset["euc_id"], self.user["user_id"])["architecture"]["components"])
        self.assertTrue(service.validation(self.asset["euc_id"], self.user["user_id"])["items"])

        unit = components["items"][0]
        overridden = service.override(self.asset["euc_id"], unit["unit_id"], self.user["user_id"],
                                      "RETAIN_IN_EXCEL", "Domain owner requires a transitional Excel surface.")
        self.assertEqual(unit["engine_mode"], overridden["engine_mode"])
        self.assertEqual("RETAIN_IN_EXCEL", overridden["effective_mode"])
        self.assertEqual("Domain owner requires a transitional Excel surface.", overridden["override_reason"])

        conn = database._get_connection()
        try:
            preview = ledger_for_connection(conn).collect_garbage(conn, dry_run=True)
            self.assertEqual(0, preview["objects_candidates"])
        finally:
            conn.close()


if __name__ == "__main__":
    unittest.main()
