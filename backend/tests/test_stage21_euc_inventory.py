import io
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from app import database
from app.euc.service import analyze_euc, export_manifest, get_inventory, ingest_euc, list_assets
from app.euc.validation import EUCValidationError, validate_euc_file


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Forecast"
    sheet.append(["Units", "Price", "Revenue"])
    sheet.append([2, 10, "=A2*B2"])
    sheet.append([3, 20, "=A3*B3"])
    sheet.sheet_state = "visible"
    hidden = workbook.create_sheet("Hidden Calc")
    hidden.sheet_state = "veryHidden"
    hidden["A1"] = "=NOW()"
    sheet.add_table(__import__("openpyxl").worksheet.table.Table(displayName="ForecastTable", ref="A1:C3"))
    validation = DataValidation(type="list", formula1='"OPEN,CLOSED"')
    sheet.add_data_validation(validation)
    validation.add("A2:A3")
    workbook.defined_names.add(DefinedName("SalesRate", attr_text="Forecast!$B$2"))
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class EUCInventoryTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_db_path = database.DB_PATH
        database.DB_PATH = Path(self.temp_dir.name) / "euc.db"
        database.initialize_product_schema()
        self.user = database.get_or_create_user("inventory-owner@example.com")
        database.create_sqlite_table_from_df("QUEUE_BOARD_EUC", pd.DataFrame([{"VALUE": 1}]))
        registered = database.register_dataset("QUEUE_BOARD_EUC", self.user["user_id"], "source.xlsx", 1, 1)
        self.repository_id = registered["repository_id"]
        self.payload = workbook_bytes()

    def tearDown(self):
        database.DB_PATH = self.original_db_path
        self.temp_dir.cleanup()

    def test_validation_rejects_fake_and_unsafe_inputs(self):
        with self.assertRaises(EUCValidationError) as error:
            validate_euc_file("fake.xlsx", b"not a zip")
        self.assertEqual("CORRUPT_OPENXML", error.exception.code)
        with self.assertRaises(EUCValidationError):
            validate_euc_file("workbook.exe", b"payload")

    def test_ingestion_deduplicates_and_analysis_is_deterministic(self):
        asset = ingest_euc(self.repository_id, "forecast.xlsx", self.payload, self.user["user_id"])
        duplicate = ingest_euc(self.repository_id, "forecast-copy.xlsx", self.payload, self.user["user_id"])
        self.assertEqual(asset["euc_id"], duplicate["euc_id"])
        self.assertTrue(duplicate["deduplicated"])

        first = analyze_euc(asset["euc_id"], self.user["user_id"])
        second = analyze_euc(asset["euc_id"], self.user["user_id"])
        self.assertIn(first["status"], {"COMPLETED", "COMPLETED_WITH_WARNINGS"})
        self.assertEqual(first["manifest_hash"], second["manifest_hash"])
        self.assertEqual(3, first["summary"]["total_formulas"])
        self.assertEqual(2, first["summary"]["unique_patterns"])

        inventory = get_inventory(asset["euc_id"], self.user["user_id"], "overview,sheets,formulas,objects,dependencies", 100, 0)
        self.assertEqual(2, len(inventory["sheets"]))
        self.assertEqual(2, len(inventory["formulas"]))
        self.assertTrue(any(item["object_type"] == "TABLE" for item in inventory["objects"]))
        self.assertTrue(any(item["object_type"] == "NAMED_RANGE" for item in inventory["objects"]))

        filename, manifest = export_manifest(asset["euc_id"], self.user["user_id"])
        self.assertTrue(filename.endswith("_inventory.json"))
        self.assertIn(b'"analysis_version": "2.1.0"', manifest)
        self.assertEqual(1, len(list_assets(self.user["user_id"])))

    def test_source_and_manifest_are_gc_roots(self):
        asset = ingest_euc(self.repository_id, "forecast.xlsx", self.payload, self.user["user_id"])
        analysis = analyze_euc(asset["euc_id"], self.user["user_id"])
        conn = database._get_connection()
        try:
            from app.services.semantic_ledger_service import ledger_for_connection
            preview = ledger_for_connection(conn).collect_garbage(conn, dry_run=True)
            roots = {row[0] for row in conn.execute("SELECT OBJECT_HASH FROM STORAGE_OBJECTS")}
            self.assertIn(asset["file_hash"], roots)
            self.assertIn(analysis["manifest_hash"], roots)
            self.assertEqual(0, preview["objects_candidates"])
        finally:
            conn.close()


if __name__ == "__main__":
    unittest.main()
