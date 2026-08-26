import io
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from app import database
from app.euc.dependency.services import DependencyService
from app.euc.intelligence.findings import FindingEngine
from app.euc.intelligence.services import IntelligenceService
from app.euc.service import analyze_euc, ingest_euc
from app.services.semantic_ledger_service import ledger_for_connection


def intelligence_workbook() -> bytes:
    workbook = Workbook()
    inputs = workbook.active
    inputs.title = "Inputs"
    inputs.append(["Value", "Choice"])
    inputs.append([10, "A"])
    validation = DataValidation(type="list", formula1='"A,B,C"')
    inputs.add_data_validation(validation)
    validation.add(inputs["B2"])
    calc = workbook.create_sheet("Calc")
    calc.append(["Result", "Dynamic", "Broken"])
    calc.append(["=Inputs!A2*2", '=INDIRECT("A"&2)', "=#REF!+1"])
    summary = workbook.create_sheet("Summary")
    summary.append(["Final"])
    summary.append(["=Calc!A2+5"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class IntelligenceServiceIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_db_path = database.DB_PATH
        database.DB_PATH = Path(self.temp_dir.name) / "intelligence.db"
        database.initialize_product_schema()
        self.user = database.get_or_create_user("intelligence-owner@example.com")
        database.create_sqlite_table_from_df("QUEUE_BOARD_INT", pd.DataFrame([{"VALUE": 1}]))
        registered = database.register_dataset(
            "QUEUE_BOARD_INT", self.user["user_id"], "source.xlsx", 1, 1
        )
        self.repository_id = registered["repository_id"]
        self.asset = ingest_euc(
            self.repository_id, "intelligence.xlsx", intelligence_workbook(), self.user["user_id"]
        )
        analyze_euc(self.asset["euc_id"], self.user["user_id"])
        DependencyService().build(self.asset["euc_id"], self.user["user_id"])

    def tearDown(self):
        database.DB_PATH = self.original_db_path
        self.temp_dir.cleanup()

    def test_scores_findings_lifecycle_determinism_and_gc(self):
        service = IntelligenceService()
        first = service.build(self.asset["euc_id"], self.user["user_id"], "FINANCIAL_MODEL")
        second = service.build(self.asset["euc_id"], self.user["user_id"], "FINANCIAL_MODEL")
        self.assertEqual(first["result_manifest_hash"], second["result_manifest_hash"])
        self.assertEqual("CURRENT", second["freshness"])
        self.assertEqual("FINANCIAL_MODEL", second["profile"])
        for score in second["scores"].values():
            self.assertGreaterEqual(score, 0)
            self.assertLessEqual(score, 100)
        self.assertLessEqual(second["scores"]["residual_risk"], second["scores"]["inherent_risk"])

        complexity = service.complexity(self.asset["euc_id"], self.user["user_id"])
        self.assertEqual(8, len(complexity["components"]))
        self.assertTrue(all(item["explanation"] and item["evidence"] is not None
                            for item in complexity["components"]))
        findings = service.findings(self.asset["euc_id"], self.user["user_id"])["items"]
        rules = {item["rule_id"] for item in findings}
        self.assertIn("BROKEN_REFERENCE", rules)
        self.assertIn("DYNAMIC_REFERENCE_HIGH_IMPACT", rules)
        finding = findings[0]
        updated = service.update_finding(
            self.asset["euc_id"], self.user["user_id"], finding["finding_id"],
            "ACKNOWLEDGED", "Owner confirmed this issue for remediation.",
        )
        self.assertEqual("ACKNOWLEDGED", updated["status"])
        self.assertEqual(1, len(updated["actions"]))

        conn = database._get_connection()
        try:
            preview = ledger_for_connection(conn).collect_garbage(conn, dry_run=True)
            self.assertEqual(0, preview["objects_candidates"])
        finally:
            conn.close()

    def test_candidate_rules_are_deduplicated_and_confidence_is_separate(self):
        item = {"node_id": "NODE_1", "sheet_id": "SHEET_1", "cell_address": "A2",
                "technical_criticality": 90, "downstream_count": 50, "sheet_spread": 2}
        features = {
            "candidates": {"broken": [item, item], "dynamic": [], "cycles": [],
                           "pattern_breaks": [], "formula_overrides": [],
                           "unresolved_external": [], "local_paths": [],
                           "change_hotspots": [], "uncontrolled_critical": []},
            "automation": {"vba_present": False},
        }
        findings = FindingEngine().evaluate(features, {})
        self.assertEqual(1, len(findings))
        self.assertEqual("CRITICAL", findings[0].severity)
        self.assertAlmostEqual(.99, findings[0].confidence)


if __name__ == "__main__":
    unittest.main()
