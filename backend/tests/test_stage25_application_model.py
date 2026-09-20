import io
import tempfile
import unittest
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from app import database
from app.euc.application_model import ApplicationModelService
from app.euc.dependency.services import DependencyService
from app.euc.intelligence.services import IntelligenceService
from app.euc.migration.services import MigrationService
from app.euc.service import analyze_euc, ingest_euc
from app.services.semantic_ledger_service import ledger_for_connection


def native_model_workbook() -> bytes:
    workbook = Workbook()
    products = workbook.active
    products.title = "Products"
    products.append(["Product ID", "Product Name", "Unit Price", "Active"])
    products.append(["P001", "Widget", 12.5, True])
    products.append(["P002", "Service", 20, True])
    table = Table(displayName="Products", ref="A1:D3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    products.add_table(table)
    pricing = workbook.create_sheet("Pricing")
    pricing.append(["Quantity", "Unit Price", "Net Amount"])
    pricing.append([2, 12.5, "=A2*B2"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class ApplicationModelIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_db_path = database.DB_PATH
        database.DB_PATH = Path(self.temp_dir.name) / "application-model.db"
        database.initialize_product_schema()
        self.user = database.get_or_create_user("application-owner@example.com")
        database.create_sqlite_table_from_df("QUEUE_BOARD_APP", pd.DataFrame([{"VALUE": 1}]))
        registered = database.register_dataset("QUEUE_BOARD_APP", self.user["user_id"], "source.xlsx", 1, 1)
        self.asset = ingest_euc(registered["repository_id"], "native.xlsx", native_model_workbook(), self.user["user_id"])
        analyze_euc(self.asset["euc_id"], self.user["user_id"])
        DependencyService().build(self.asset["euc_id"], self.user["user_id"])
        IntelligenceService().build(self.asset["euc_id"], self.user["user_id"])
        MigrationService().analyze(self.asset["euc_id"], self.user["user_id"])

    def tearDown(self):
        database.DB_PATH = self.original_db_path
        self.temp_dir.cleanup()

    def test_air_is_deterministic_reviewed_and_compiled_to_cas(self):
        service = ApplicationModelService()
        first = service.build(self.asset["euc_id"], self.user["user_id"])
        second = service.build(self.asset["euc_id"], self.user["user_id"])
        self.assertEqual(first["manifest_hash"], second["manifest_hash"])
        self.assertEqual("2.5.0", first["air_version"])
        self.assertGreater(first["summary"]["counts"]["entities"], 0)
        self.assertGreater(first["summary"]["counts"]["apis"], 0)
        manifest = service.manifest(self.asset["euc_id"], self.user["user_id"])
        self.assertEqual("gitwalk.air.v1", manifest["schema"])
        self.assertTrue(manifest["lineage"])
        self.assertTrue(all(item.get("provenance") for item in manifest["entities"]))

        for item in service.components(self.asset["euc_id"], self.user["user_id"])["items"]:
            if item["review_state"] == "REVIEW_REQUIRED":
                service.review(self.asset["euc_id"], item["component_id"], self.user["user_id"],
                               "CONFIRMED", "Domain owner confirmed this inferred component.")
        approved = service.approve(self.asset["euc_id"], self.user["user_id"])
        self.assertEqual("APPROVED", approved["status"])

        generated = service.generate(self.asset["euc_id"], self.user["user_id"], "SCAFFOLD")
        self.assertEqual("COMPLETED", generated["status"])
        self.assertGreater(generated["files_generated"], 10)
        filename, payload = service.download(self.asset["euc_id"], generated["generation_run_id"], self.user["user_id"])
        self.assertTrue(filename.endswith(".zip"))
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            names = set(archive.namelist())
            self.assertIn("air/application-model.json", names)
            self.assertIn("database/schema/001_initial.sql", names)
            self.assertIn("backend/app/main.py", names)
            self.assertIn("frontend/src/App.tsx", names)
            self.assertIn("frontend/index.html", names)
            for path in sorted(name for name in names if name.endswith(".py")):
                compile(archive.read(path), path, "exec")

        conn = database._get_connection()
        try:
            preview = ledger_for_connection(conn).collect_garbage(conn, dry_run=True)
            self.assertEqual(0, preview["objects_candidates"])
        finally:
            conn.close()


if __name__ == "__main__":
    unittest.main()
