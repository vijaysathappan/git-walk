import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app import database
from app.repositories.commit_store import commit_semantic_delta
from app.repositories.commit_store import branch_change_timeline
from app.repositories.governance_store import workbook_change_activity
from app.services.workbook_service import _write_branch_xlsx


class MultiSheetRepositoryTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_db_path = database.DB_PATH
        database.DB_PATH = Path(self.temp_dir.name) / "multisheet.db"
        database.initialize_product_schema()
        self.owner = database.get_or_create_user("owner@example.com")
        database.create_sqlite_table_from_df(
            "QUEUE_BOARD_MULTI",
            pd.DataFrame([{"SKU": "A-1", "AMOUNT": 10}]),
        )
        database.create_sqlite_table_from_df(
            "SHEET_DATA_LOOKUPS",
            pd.DataFrame([{"CODE": "IN", "LABEL": "India"}]),
        )
        self.repository = database.register_dataset(
            "QUEUE_BOARD_MULTI",
            self.owner["user_id"],
            "finance.xlsx",
            1,
            2,
            repository_name="finance-controls",
            sheet_tables=[
                {"name": "Transactions", "table_id": "QUEUE_BOARD_MULTI"},
                {"name": "Lookups", "table_id": "SHEET_DATA_LOOKUPS"},
            ],
        )
        database.store_initial_formula_metadata(
            "QUEUE_BOARD_MULTI",
            {"Transactions": [{
                "row_position": 0, "column_position": 1, "formula": "=5+5",
            }]},
        )
        self.copy = database.create_working_copy(
            "QUEUE_BOARD_MULTI", self.owner["user_id"], self.owner["email"]
        )

    def tearDown(self):
        database.DB_PATH = self.original_db_path
        self.temp_dir.cleanup()

    def _snapshot(self):
        return database.get_table_snapshot(self.copy["table_id"])

    def test_signed_checkout_can_resume_or_create_a_new_branch(self):
        authenticated = database.authenticate_working_copy_identity(
            self.copy["table_id"], self.copy["repository_id"],
            self.copy["branch_id"], self.copy["working_copy_id"],
            self.copy["base_commit_id"], self.copy["issued_at"],
            self.copy["signature"],
        )
        self.assertEqual(self.owner["user_id"], authenticated["user_id"])

        options = database.working_copy_checkout_options(
            "QUEUE_BOARD_MULTI", self.owner["user_id"]
        )
        self.assertTrue(options["can_edit"])
        self.assertEqual(self.copy["branch_id"], options["branches"][0]["branch_id"])

        resumed = database.create_working_copy(
            "QUEUE_BOARD_MULTI", self.owner["user_id"], self.owner["email"],
            branch_mode="continue", branch_id=self.copy["branch_id"],
        )
        created = database.create_working_copy(
            "QUEUE_BOARD_MULTI", self.owner["user_id"], self.owner["email"],
            branch_mode="new",
        )
        self.assertEqual(self.copy["branch_id"], resumed["branch_id"])
        self.assertNotEqual(self.copy["branch_id"], created["branch_id"])

    def test_activity_and_merge_timeline_include_only_committed_mutations(self):
        snapshot = self._snapshot()
        sheet = snapshot["semantic"]["sheets"][0]
        row = sheet["rows"][0]
        column = next(item for item in sheet["columns"] if item["name"] == "AMOUNT")
        result = commit_semantic_delta(
            table_id=self.copy["table_id"],
            repository_id=self.copy["repository_id"],
            branch_id=self.copy["branch_id"],
            expected_head_commit_id=snapshot["head_commit_id"],
            base_version=snapshot["version"],
            changes=[{
                "operation_type": "CELL_VALUE_UPDATE",
                "sheet_id": sheet["sheet_id"],
                "row_id": row["row_id"],
                "column_id": column["column_id"],
                "new_value": 25,
            }],
            user_id=self.owner["user_id"],
            user_email=self.owner["email"],
            message="Adjust amount",
        )

        activity = workbook_change_activity(self.copy["branch_id"])
        self.assertEqual(1, activity["event_count"])
        self.assertEqual("updated", activity["events"][0]["action"])
        timeline = branch_change_timeline(
            self.copy["branch_id"], result["commit_id"], self.copy["base_commit_id"]
        )
        self.assertEqual(["CELL_VALUE_UPDATE"], [item["operation_type"] for item in timeline])
        self.assertEqual("Adjust amount", timeline[0]["commit_message"])

    def test_secondary_sheet_commit_and_complete_export(self):
        snapshot = self._snapshot()
        self.assertEqual(["Transactions", "Lookups"], [
            sheet["name"] for sheet in snapshot["semantic"]["sheets"]
        ])
        lookup = snapshot["semantic"]["sheets"][1]
        label = next(column for column in lookup["columns"] if column["name"] == "LABEL")
        row = lookup["rows"][0]

        result = commit_semantic_delta(
            table_id=self.copy["table_id"],
            repository_id=self.copy["repository_id"],
            branch_id=self.copy["branch_id"],
            expected_head_commit_id=snapshot["head_commit_id"],
            base_version=snapshot["version"],
            changes=[{
                "operation_type": "CELL_VALUE_UPDATE",
                "sheet_id": lookup["sheet_id"],
                "row_id": row["row_id"],
                "column_id": label["column_id"],
                "new_value": "India updated",
            }],
            user_id=self.owner["user_id"],
            user_email=self.owner["email"],
            message="Update lookup label",
        )
        self.assertEqual(1, result["version"])

        connection = database._get_connection()
        try:
            mapped_table = connection.execute(
                "SELECT DATA_TABLE_ID FROM BRANCH_SHEET_TABLES WHERE BRANCH_ID=? AND SHEET_ID=?",
                (self.copy["branch_id"], lookup["sheet_id"]),
            ).fetchone()[0]
        finally:
            connection.close()
        self.assertEqual(
            "India updated", database.read_cell(mapped_table, "LABEL", 1)
        )

        latest = self._snapshot()
        new_sheet_id = "SHEET_CLIENTNOTES"
        commit_semantic_delta(
            table_id=self.copy["table_id"],
            repository_id=self.copy["repository_id"],
            branch_id=self.copy["branch_id"],
            expected_head_commit_id=latest["head_commit_id"],
            base_version=latest["version"],
            changes=[
                {
                    "operation_type": "SHEET_CREATE", "sheet_id": new_sheet_id,
                    "new_value": "Notes", "new_row_position": 2,
                },
                {
                    "operation_type": "COLUMN_INSERT", "sheet_id": new_sheet_id,
                    "column_id": "COL_NOTE", "new_value": "NOTE",
                    "new_column_position": 0, "new_data_type": "TEXT",
                },
                {
                    "operation_type": "ROW_INSERT", "sheet_id": new_sheet_id,
                    "row_id": "ROW_NOTE1", "new_row_position": 0,
                    "new_value": {"COL_NOTE": "Reviewed"},
                },
            ],
            user_id=self.owner["user_id"],
            user_email=self.owner["email"],
            message="Add review notes sheet",
        )
        final_snapshot = self._snapshot()
        self.assertEqual(3, len(final_snapshot["semantic"]["sheets"]))

        output = Path(self.temp_dir.name) / "branch.xlsx"
        _write_branch_xlsx(self.copy["table_id"], output)
        workbook = load_workbook(output, data_only=False)
        try:
            self.assertEqual(["Transactions", "Lookups", "Notes"], workbook.sheetnames)
            self.assertEqual("=5+5", workbook["Transactions"]["B2"].value)
            self.assertEqual("India updated", workbook["Lookups"]["B2"].value)
            self.assertEqual("Reviewed", workbook["Notes"]["A2"].value)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
