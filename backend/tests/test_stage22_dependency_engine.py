import io
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from app import database
from app.euc.dependency.graph import DependencyGraphBuilder, ResolverCatalog
from app.euc.dependency.graph.cycles import find_cycles
from app.euc.dependency.parser import FormulaParser
from app.euc.dependency.resolver import parse_reference
from app.euc.dependency.services import DependencyService
from app.euc.service import analyze_euc, ingest_euc


def catalog() -> ResolverCatalog:
    return ResolverCatalog(
        repository_id="REPO_TEST",
        euc_id="EUC_TEST",
        analysis_id="AN_TEST",
        sheets_by_name={"inputs": "SHEET_INPUTS", "calc": "SHEET_CALC", "summary": "SHEET_SUMMARY"},
        sheet_names={"SHEET_INPUTS": "Inputs", "SHEET_CALC": "Calc", "SHEET_SUMMARY": "Summary"},
        sheet_dimensions={"SHEET_INPUTS": (100000, 10), "SHEET_CALC": (100, 10), "SHEET_SUMMARY": (100, 10)},
    )


def occurrence(sheet_id: str, cell: str, formula: str, position: int) -> dict:
    return {
        "pattern_id": f"PAT_{position}", "sheet_id": sheet_id,
        "cell_address": cell, "raw_formula": formula,
        "normalized_hash": f"HASH_{position}",
    }


def workbook_bytes() -> bytes:
    workbook = Workbook()
    inputs = workbook.active
    inputs.title = "Inputs"
    inputs.append(["Value", "Label"])
    inputs.append([10, "Base"])
    calc = workbook.create_sheet("Calc")
    calc.append(["Result", "Adjusted"])
    calc.append(["=Inputs!A2*2", "=A2+5"])
    summary = workbook.create_sheet("Summary")
    summary.append(["Final"])
    summary.append(["=Calc!B2/3"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class FormulaDependencyUnitTests(unittest.TestCase):
    def test_reference_types_preserve_excel_semantics(self):
        absolute = parse_reference("Assumptions!$B$5")
        self.assertEqual("CELL", absolute.kind)
        self.assertEqual("Assumptions", absolute.sheet_name)
        self.assertTrue(absolute.absolute_column)
        self.assertTrue(absolute.absolute_row)
        self.assertEqual("COLUMN_RANGE", parse_reference("Sales!$B:$D").kind)
        self.assertEqual("ROW_RANGE", parse_reference("1:3").kind)
        self.assertEqual("TABLE_COLUMN", parse_reference("Sales[Revenue]").kind)
        self.assertEqual("Revenue", parse_reference("Sales[Revenue]").table_column)
        self.assertEqual("NAMED_RANGE", parse_reference("TAX_RATE").kind)
        external = parse_reference("'[Pricing.xlsx]Rates'!$D$5")
        self.assertEqual("Pricing.xlsx", external.workbook_name)
        self.assertEqual("Rates", external.sheet_name)

    def test_parser_preserves_function_operator_and_references(self):
        ast = FormulaParser().parse("=IF(A1>0,SUM(B1:B3)*2,-C1)")
        nodes = list(ast.walk())
        self.assertEqual("function", ast.kind)
        self.assertEqual("IF", ast.value)
        self.assertEqual({"A1", "B1:B3", "C1"}, {node.value for node in nodes if node.kind == "reference"})
        self.assertIn("*", {node.value for node in nodes if node.kind == "binary_operator"})

    def test_cross_sheet_chain_has_correct_upstream_and_downstream_direction(self):
        graph, _, _ = DependencyGraphBuilder(catalog()).build([
            occurrence("SHEET_CALC", "A2", "=Inputs!A2*2", 1),
            occurrence("SHEET_CALC", "B2", "=A2+5", 2),
            occurrence("SHEET_SUMMARY", "A2", "=Calc!B2/3", 3),
        ])
        upstream, downstream = graph.adjacency()
        calc_a2 = catalog().cell_node("SHEET_CALC", "A2").node_id
        calc_b2 = catalog().cell_node("SHEET_CALC", "B2").node_id
        summary_a2 = catalog().cell_node("SHEET_SUMMARY", "A2").node_id
        input_a2 = catalog().cell_node("SHEET_INPUTS", "A2").node_id
        self.assertIn(input_a2, upstream[calc_a2])
        self.assertIn(calc_b2, downstream[calc_a2])
        self.assertIn(summary_a2, downstream[calc_b2])

    def test_cycle_dynamic_broken_and_large_range_are_explicit(self):
        graph, _, _ = DependencyGraphBuilder(catalog()).build([
            occurrence("SHEET_CALC", "A1", "=B1+1", 1),
            occurrence("SHEET_CALC", "B1", "=C1+1", 2),
            occurrence("SHEET_CALC", "C1", "=A1+1", 3),
            occurrence("SHEET_CALC", "D1", "=INDIRECT(\"A\"&1)", 4),
            occurrence("SHEET_CALC", "E1", "=#REF!+1", 5),
            occurrence("SHEET_SUMMARY", "A1", "=SUM(Inputs!A1:A100000)", 6),
        ])
        upstream, _ = graph.adjacency()
        self.assertEqual(1, len(find_cycles(upstream)))
        self.assertTrue(any(edge.category == "DYNAMIC" and edge.resolution_status == "AMBIGUOUS" for edge in graph.edges.values()))
        self.assertTrue(any(edge.resolution_status == "BROKEN" for edge in graph.edges.values()))
        large = [edge for edge in graph.edges.values() if edge.logical_cardinality == 100000]
        self.assertEqual(1, len(large))


class DependencyServiceIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_db_path = database.DB_PATH
        database.DB_PATH = Path(self.temp_dir.name) / "dependency.db"
        database.initialize_product_schema()
        self.user = database.get_or_create_user("dependency-owner@example.com")
        database.create_sqlite_table_from_df("QUEUE_BOARD_DEP", pd.DataFrame([{"VALUE": 1}]))
        registered = database.register_dataset("QUEUE_BOARD_DEP", self.user["user_id"], "source.xlsx", 1, 1)
        self.repository_id = registered["repository_id"]

    def tearDown(self):
        database.DB_PATH = self.original_db_path
        self.temp_dir.cleanup()

    def test_build_is_deterministic_queryable_and_gc_rooted(self):
        asset = ingest_euc(self.repository_id, "dependency.xlsx", workbook_bytes(), self.user["user_id"])
        analyze_euc(asset["euc_id"], self.user["user_id"])
        service = DependencyService()
        first = service.build(asset["euc_id"], self.user["user_id"])
        second = service.build(asset["euc_id"], self.user["user_id"])
        self.assertEqual(first["graph_manifest_hash"], second["graph_manifest_hash"])
        self.assertEqual("CURRENT", second["freshness"])
        self.assertEqual(3, second["metrics"]["formula_nodes"])
        self.assertEqual(100.0, second["metrics"]["dependency_coverage"])

        sheets = service.sheets(asset["euc_id"], self.user["user_id"])
        calc = next(sheet for sheet in sheets["nodes"] if sheet["name"] == "Calc")
        node = service.resolve_cell_node(asset["euc_id"], self.user["user_id"], calc["sheet_id"], "A2")
        lineage = service.lineage(asset["euc_id"], self.user["user_id"], node["node_id"], "both", 6)
        self.assertEqual(1, len(lineage["directions"]["upstream"]["distances"]))
        self.assertEqual(2, len(lineage["directions"]["downstream"]["distances"]))
        impact = service.impact(asset["euc_id"], self.user["user_id"], node["node_id"], 6)
        self.assertEqual(2, impact["total_downstream"])

        conn = database._get_connection()
        try:
            from app.services.semantic_ledger_service import ledger_for_connection
            preview = ledger_for_connection(conn).collect_garbage(conn, dry_run=True)
            self.assertEqual(0, preview["objects_candidates"])
        finally:
            conn.close()


if __name__ == "__main__":
    unittest.main()
