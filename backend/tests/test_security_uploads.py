import re
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.security_uploads import worksheet_data_dimensions


class DimensionlessWorksheet:
    def __init__(self, resolved_rows=None, resolved_columns=None):
        self.max_row = None
        self.max_column = None
        self.resolved_rows = resolved_rows
        self.resolved_columns = resolved_columns
        self.force_used = False

    def calculate_dimension(self, force=False):
        self.force_used = force
        self.max_row = self.resolved_rows
        self.max_column = self.resolved_columns
        return "A1"


class WorksheetDimensionTests(unittest.TestCase):
    def test_recalculates_missing_read_only_dimensions(self):
        sheet = DimensionlessWorksheet(resolved_rows=31, resolved_columns=7)

        self.assertEqual(worksheet_data_dimensions(sheet), (30, 7))
        self.assertTrue(sheet.force_used)

    def test_empty_dimensionless_sheet_is_zero_sized(self):
        sheet = DimensionlessWorksheet()

        self.assertEqual(worksheet_data_dimensions(sheet), (0, 0))

    def test_reads_xlsx_when_dimension_metadata_is_missing(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            original = Path(temp_dir) / "original.xlsx"
            dimensionless = Path(temp_dir) / "dimensionless.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["PRODUCT_ID", "STATUS"])
            sheet.append(["P001", "Active"])
            workbook.save(original)

            with zipfile.ZipFile(original) as source, zipfile.ZipFile(
                dimensionless, "w"
            ) as target:
                for item in source.infolist():
                    data = source.read(item.filename)
                    if item.filename == "xl/worksheets/sheet1.xml":
                        data = re.sub(rb"<dimension[^>]*/>", b"", data)
                    target.writestr(item, data)

            uploaded = load_workbook(dimensionless, read_only=True)
            try:
                self.assertIsNone(uploaded.active.max_row)
                self.assertEqual(worksheet_data_dimensions(uploaded.active), (1, 2))
                self.assertEqual(
                    list(uploaded.active.iter_rows(values_only=True))[1],
                    ("P001", "Active"),
                )
            finally:
                uploaded.close()


if __name__ == "__main__":
    unittest.main()
