"""Build reusable Stage 2.3 feature indexes from inventory, graph, and history."""

from __future__ import annotations

import io
import json
import re
from collections import Counter
from typing import Any

from openpyxl import load_workbook

from ....config import settings
from ....services.semantic_ledger_service import ledger_for_connection


LOCAL_PATH = re.compile(r"(?:[A-Za-z]:\\|C:/Users/|\\\\[^\\]+\\)", re.I)


class FeatureBuilder:
    def __init__(self, conn, asset, inventory_run, dependency_run):
        self.conn = conn
        self.asset = asset
        self.inventory_run = inventory_run
        self.dependency_run = dependency_run
        self.analysis_id = inventory_run["ANALYSIS_ID"]
        self.dependency_run_id = dependency_run["DEPENDENCY_RUN_ID"]

    def build(self) -> dict[str, Any]:
        inventory = json.loads(self.inventory_run["SUMMARY_JSON"] or "{}")
        dependency = json.loads(self.dependency_run["METRICS_JSON"] or "{}")
        sheets = [self._lower(row) for row in self.conn.execute(
            "SELECT * FROM EUC_SHEET_INVENTORY WHERE ANALYSIS_ID=? ORDER BY SHEET_POSITION",
            (self.analysis_id,),
        )]
        object_counts = dict(Counter(row[0] for row in self.conn.execute(
            "SELECT OBJECT_TYPE FROM EUC_OBJECT_INVENTORY WHERE ANALYSIS_ID=?", (self.analysis_id,)
        )))
        ast = self._ast_features()
        history = self._history_features()
        critical_nodes = [self._lower(row) for row in self.conn.execute(
            """SELECT * FROM EUC_DEPENDENCY_NODES WHERE DEPENDENCY_RUN_ID=?
               AND TECHNICAL_CRITICALITY>=? ORDER BY TECHNICAL_CRITICALITY DESC""",
            (self.dependency_run_id, settings.intelligence_criticality_threshold),
        )]
        controls = self._control_features(sheets, critical_nodes, dependency)
        candidates = self._candidate_indexes(critical_nodes, controls)
        data = self._data_features(sheets)
        external = {
            "external_links": inventory.get("external_links", 0),
            "connections": inventory.get("connections", 0),
            "unresolved": len(candidates["unresolved_external"]),
            "local_paths": len(candidates["local_paths"]),
            "credentials_present": bool(self.conn.execute(
                "SELECT 1 FROM EUC_CONNECTIONS WHERE ANALYSIS_ID=? AND CREDENTIAL_PRESENT=1 LIMIT 1",
                (self.analysis_id,),
            ).fetchone()),
        }
        automation = {
            "vba_present": bool(inventory.get("vba_present") or inventory.get("macro_enabled")),
            "power_queries": inventory.get("power_queries", 0),
            "connections": inventory.get("connections", 0),
        }
        formula_errors = self.conn.execute(
            """SELECT COUNT(*) FROM EUC_FORMULA_OCCURRENCES WHERE ANALYSIS_ID=? AND
               (RAW_FORMULA LIKE '%#REF!%' OR RAW_FORMULA LIKE '%#DIV/0!%' OR
                RAW_FORMULA LIKE '%#VALUE!%' OR RAW_FORMULA LIKE '%#NAME?%')""",
            (self.analysis_id,),
        ).fetchone()[0]
        return {
            "inventory": inventory, "dependency": dependency, "sheets": sheets,
            "objects": object_counts, "ast": ast, "history": history,
            "data": data, "external": external, "automation": automation,
            "controls": controls, "candidates": candidates,
            "candidate_counts": {key: len(value) for key, value in candidates.items()},
            "formula_errors": formula_errors, "critical_nodes": len(critical_nodes),
            "change_hotspot_count": len(candidates["change_hotspots"]),
        }

    def _ast_features(self) -> dict:
        row = self.conn.execute(
            """SELECT COUNT(*),COALESCE(AVG(AST_DEPTH),0),COALESCE(MAX(AST_DEPTH),0),
                      COALESCE(SUM(FUNCTION_COUNT),0),COALESCE(SUM(REFERENCE_COUNT),0),
                      COALESCE(SUM(RANGE_COUNT),0),COALESCE(SUM(DYNAMIC_REFERENCE_COUNT),0)
               FROM EUC_FORMULA_ASTS WHERE DEPENDENCY_RUN_ID=?""",
            (self.dependency_run_id,),
        ).fetchone()
        functions = set()
        for item in self.conn.execute("SELECT FUNCTIONS_JSON FROM EUC_FORMULA_PATTERNS WHERE ANALYSIS_ID=?", (self.analysis_id,)):
            functions.update(json.loads(item[0] or "[]"))
        return {"patterns": row[0], "average_depth": round(row[1], 2), "maximum_depth": row[2],
                "function_count": row[3], "reference_count": row[4], "range_count": row[5],
                "dynamic_reference_count": row[6], "function_diversity": len(functions),
                "functions": sorted(functions)}

    def _history_features(self) -> dict:
        repository_id = self.asset["REPOSITORY_ID"]
        commits = self.conn.execute(
            """SELECT COUNT(*),COUNT(DISTINCT AUTHOR_USER_ID),
                      COALESCE(SUM(CASE WHEN REVERTS_COMMIT_ID IS NOT NULL THEN 1 ELSE 0 END),0)
               FROM COMMITS WHERE REPOSITORY_ID=?""", (repository_id,)
        ).fetchone()
        changes = self.conn.execute(
            """SELECT COUNT(*),COALESCE(SUM(CASE WHEN OLD_FORMULA IS NOT NEW_FORMULA THEN 1 ELSE 0 END),0)
               FROM COMMIT_CHANGES WHERE REPOSITORY_ID=?""", (repository_id,)
        ).fetchone()
        conflicts = self.conn.execute(
            """SELECT COUNT(*) FROM MERGE_CONFLICTS C JOIN MERGE_REQUESTS M
               ON M.MERGE_REQUEST_ID=C.MERGE_REQUEST_ID WHERE M.REPOSITORY_ID=?""", (repository_id,)
        ).fetchone()[0]
        return {"commits": commits[0], "contributors": commits[1], "reversions": commits[2],
                "changes": changes[0], "formula_changes": changes[1], "conflicts": conflicts}

    def _data_features(self, sheets: list[dict]) -> dict:
        mixed = self.conn.execute(
            """SELECT COUNT(*) FROM SHEET_COLUMNS C JOIN WORKBOOK_REPOSITORIES R
               ON R.DEFAULT_BRANCH_ID=C.BRANCH_ID WHERE R.REPOSITORY_ID=? AND C.DATA_TYPE='MIXED'""",
            (self.asset["REPOSITORY_ID"],),
        ).fetchone()[0]
        sparse = sum(
            (sheet.get("max_row") or 0) * (sheet.get("max_column") or 0) > max(1, sheet.get("used_cell_count") or 0) * 4
            for sheet in sheets
        )
        return {"mixed_type_columns": mixed, "sparse_sheets": sparse,
                "blank_styled_cells": sum(sheet.get("blank_styled_cell_count") or 0 for sheet in sheets)}

    def _control_features(self, sheets: list[dict], critical_nodes: list[dict], dependency: dict) -> dict:
        protected_sheets: set[str] = set()
        locked_nodes: set[str] = set()
        validated_nodes: set[str] = set()
        validations = sum(sheet.get("validation_count") or 0 for sheet in sheets)
        if self.asset["FILE_TYPE"] != "csv":
            payload = ledger_for_connection(self.conn).objects.get_bytes(self.conn, self.asset["ORIGINAL_OBJECT_HASH"])
            workbook = load_workbook(io.BytesIO(payload), data_only=False, keep_links=False,
                                     keep_vba=self.asset["FILE_TYPE"] == "xlsm")
            try:
                sheet_names = {sheet["sheet_id"]: sheet["sheet_name"] for sheet in sheets}
                for node in critical_nodes:
                    sheet_name = sheet_names.get(node.get("sheet_id"))
                    address = node.get("cell_address")
                    if not sheet_name or not address or sheet_name not in workbook.sheetnames:
                        continue
                    worksheet = workbook[sheet_name]
                    if worksheet.protection.sheet:
                        protected_sheets.add(node["sheet_id"])
                        if worksheet[address].protection.locked:
                            locked_nodes.add(node["node_id"])
                    if any(address in validation.sqref for validation in worksheet.data_validations.dataValidation):
                        validated_nodes.add(node["node_id"])
            finally:
                workbook.close()
        controlled = locked_nodes | validated_nodes
        criticality = {item["node_id"]: float(item.get("technical_criticality") or 0) for item in critical_nodes}
        repository = self.conn.execute(
            "SELECT MAIN_PROTECTED,DEFAULT_BRANCH_ID FROM WORKBOOK_REPOSITORIES WHERE REPOSITORY_ID=?",
            (self.asset["REPOSITORY_ID"],),
        ).fetchone()
        governance = {
            "audit_ledger": True,
            "version_history": bool(self.conn.execute("SELECT 1 FROM COMMITS WHERE REPOSITORY_ID=? LIMIT 1", (self.asset["REPOSITORY_ID"],)).fetchone()),
            "review_workflow": True,
            "branch_protection": bool(repository["MAIN_PROTECTED"]),
            "rollback": True,
            "change_attribution": True,
        }
        return {
            "validations": validations, "protected_sheets": len(protected_sheets),
            "protected_formula_nodes": len(locked_nodes), "validated_critical_nodes": len(validated_nodes),
            "controlled_critical_nodes": len(controlled), "critical_nodes": len(critical_nodes),
            "formula_nodes": dependency.get("formula_nodes", 0),
            "controlled_criticality": sum(criticality.get(node, 0) for node in controlled),
            "total_criticality": sum(criticality.values()), "governance": governance,
            "controlled_node_ids": sorted(controlled),
        }

    def _candidate_indexes(self, critical_nodes: list[dict], controls: dict) -> dict[str, list[dict]]:
        broken = [self._lower(row) for row in self.conn.execute(
            """SELECT E.*,N.DISPLAY_NAME AS SOURCE_NAME,N.SHEET_ID,N.CELL_ADDRESS,
                      N.DOWNSTREAM_COUNT,N.SHEET_SPREAD,N.TECHNICAL_CRITICALITY
               FROM EUC_DEPENDENCY_EDGES E JOIN EUC_DEPENDENCY_NODES N
                 ON N.DEPENDENCY_RUN_ID=E.DEPENDENCY_RUN_ID AND N.NODE_ID=E.SOURCE_NODE_ID
               WHERE E.DEPENDENCY_RUN_ID=? AND E.RESOLUTION_STATUS='BROKEN' LIMIT ?""",
            (self.dependency_run_id, settings.intelligence_finding_limit),
        )]
        dynamic = [self._lower(row) for row in self.conn.execute(
            """SELECT E.*,N.DISPLAY_NAME AS SOURCE_NAME,N.SHEET_ID,N.CELL_ADDRESS,
                      N.DOWNSTREAM_COUNT,N.SHEET_SPREAD,N.TECHNICAL_CRITICALITY
               FROM EUC_DEPENDENCY_EDGES E JOIN EUC_DEPENDENCY_NODES N
                 ON N.DEPENDENCY_RUN_ID=E.DEPENDENCY_RUN_ID AND N.NODE_ID=E.SOURCE_NODE_ID
               WHERE E.DEPENDENCY_RUN_ID=? AND E.DEPENDENCY_CATEGORY='DYNAMIC' LIMIT ?""",
            (self.dependency_run_id, settings.intelligence_finding_limit),
        )]
        cycles = [dict(self._lower(row), sheets=json.loads(row["SHEETS_JSON"] or "[]")) for row in self.conn.execute(
            "SELECT * FROM EUC_DEPENDENCY_CYCLES WHERE DEPENDENCY_RUN_ID=?", (self.dependency_run_id,)
        )]
        external_rows = [self._lower(row) for row in self.conn.execute(
            "SELECT * FROM EUC_EXTERNAL_LINKS WHERE ANALYSIS_ID=?", (self.analysis_id,)
        )]
        unresolved = [item for item in external_rows if item.get("resolution_status") != "RESOLVED"]
        local_paths = [{"path": value, "source": item} for item in external_rows
                       for value in (item.get("source_euc_reference"), item.get("source_address"))
                       if value and LOCAL_PATH.search(value)]
        pattern_breaks = self._pattern_breaks()
        node_map = {(item["sheet_id"], item["cell_address"]): item for item in critical_nodes}
        for item in pattern_breaks:
            item.update(node_map.get((item.get("sheet_id"), item.get("cell_address")), {}))
        overrides = self._formula_overrides()
        hotspots = self._change_hotspots()
        controlled = set(controls["controlled_node_ids"])
        uncontrolled = [item for item in critical_nodes if item["node_id"] not in controlled and item.get("metadata", {}).get("formula")]
        return {"broken": broken, "dynamic": dynamic, "cycles": cycles,
                "pattern_breaks": pattern_breaks, "formula_overrides": overrides,
                "unresolved_external": unresolved, "local_paths": local_paths,
                "change_hotspots": hotspots, "uncontrolled_critical": uncontrolled}

    def _pattern_breaks(self) -> list[dict]:
        manifest_hash = self.dependency_run["GRAPH_MANIFEST_HASH"]
        if not manifest_hash:
            return []
        ledger = ledger_for_connection(self.conn)
        manifest = ledger.objects.get(self.conn, manifest_hash)
        object_hash = manifest.get("pattern_breaks_hash")
        return ledger.objects.get(self.conn, object_hash) if object_hash else []

    def _formula_overrides(self) -> list[dict]:
        rows = self.conn.execute(
            """SELECT C.*,N.NODE_ID,N.CELL_ADDRESS,N.DOWNSTREAM_COUNT,N.SHEET_SPREAD,N.TECHNICAL_CRITICALITY
               FROM COMMIT_CHANGES C LEFT JOIN EUC_DEPENDENCY_NODES N
                 ON N.DEPENDENCY_RUN_ID=? AND N.SHEET_ID=C.SHEET_ID
                AND N.ROW_ID=C.ROW_ID AND N.COLUMN_ID=C.COLUMN_ID
               WHERE C.REPOSITORY_ID=? AND C.OLD_FORMULA IS NOT NULL
                 AND C.NEW_FORMULA IS NULL AND C.NEW_VALUE IS NOT NULL
               ORDER BY C.CREATED_AT DESC LIMIT ?""",
            (self.dependency_run_id, self.asset["REPOSITORY_ID"], settings.intelligence_finding_limit),
        )
        unique = {}
        for row in rows:
            item = self._lower(row)
            key = (item.get("sheet_id"), item.get("row_id"), item.get("column_id"))
            item["confidence"] = .98 if item.get("old_formula") else .75
            unique.setdefault(key, item)
        return list(unique.values())

    def _change_hotspots(self) -> list[dict]:
        rows = self.conn.execute(
            """SELECT C.SHEET_ID,C.ROW_ID,C.COLUMN_ID,COUNT(*) AS CHANGE_FREQUENCY,
                      COUNT(DISTINCT C.COMMIT_ID) AS COMMIT_FREQUENCY,N.*
               FROM COMMIT_CHANGES C JOIN EUC_DEPENDENCY_NODES N
                 ON N.DEPENDENCY_RUN_ID=? AND N.SHEET_ID=C.SHEET_ID
                AND N.ROW_ID=C.ROW_ID AND N.COLUMN_ID=C.COLUMN_ID
               WHERE C.REPOSITORY_ID=? GROUP BY C.SHEET_ID,C.ROW_ID,C.COLUMN_ID
               HAVING COUNT(*)>=3 ORDER BY COUNT(*)*N.TECHNICAL_CRITICALITY DESC LIMIT 200""",
            (self.dependency_run_id, self.asset["REPOSITORY_ID"]),
        )
        return [dict(self._lower(row), hotspot_score=round(min(100, row["CHANGE_FREQUENCY"] * row["TECHNICAL_CRITICALITY"] / 10), 2)) for row in rows]

    @staticmethod
    def _lower(row) -> dict:
        result = {key.lower(): row[key] for key in row.keys()}
        if "metadata_json" in result:
            result["metadata"] = json.loads(result.pop("metadata_json") or "{}")
        return result
