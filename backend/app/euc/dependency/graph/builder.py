"""Resolve Stage 2.1 formula occurrences into a semantic dependency graph."""

from __future__ import annotations

import hashlib
import json
from collections import defaultdict
from dataclasses import dataclass, field
from typing import Any

from openpyxl.utils import column_index_from_string, coordinate_to_tuple, get_column_letter, range_boundaries

from ....config import settings
from ..parser import FormulaParseError, FormulaParser
from ..resolver import FormulaReference, parse_reference
from .models import DependencyEdge, DependencyGraph, DependencyNode


def _safe(value: str) -> str:
    return value.replace(":", "_").replace(" ", "_").upper()


def _cell_address(value: str) -> str:
    return value.replace("$", "").upper()


@dataclass
class ResolverCatalog:
    repository_id: str
    euc_id: str
    analysis_id: str
    sheets_by_name: dict[str, str]
    sheet_names: dict[str, str]
    sheet_dimensions: dict[str, tuple[int, int]]
    stable_rows: dict[tuple[str, int], str] = field(default_factory=dict)
    stable_columns: dict[tuple[str, int], str] = field(default_factory=dict)
    named_ranges: dict[str, dict[str, Any]] = field(default_factory=dict)
    tables: dict[tuple[str, str], dict[str, Any]] = field(default_factory=dict)
    external_assets: dict[str, list[dict[str, str]]] = field(default_factory=dict)

    def sheet_id(self, name: str | None, fallback: str) -> str | None:
        return fallback if not name else self.sheets_by_name.get(name.casefold())

    def cell_node(self, sheet_id: str, address: str, **metadata: Any) -> DependencyNode:
        row, column = coordinate_to_tuple(_cell_address(address))
        row_id = self.stable_rows.get((sheet_id, row), f"ROWPOS_{row}")
        column_id = self.stable_columns.get((sheet_id, column), f"COLPOS_{column}")
        node_id = f"CELL:{self.repository_id}:{sheet_id}:{row_id}:{column_id}"
        return DependencyNode(node_id, "CELL", f"{self.sheet_names.get(sheet_id, sheet_id)}!{get_column_letter(column)}{row}",
                              sheet_id, row_id, column_id, f"{get_column_letter(column)}{row}", metadata)


class DependencyGraphBuilder:
    version = "dependency-builder@1.0"

    def __init__(self, catalog: ResolverCatalog):
        self.catalog = catalog
        self.parser = FormulaParser()
        self.graph = DependencyGraph()
        self.ast_by_pattern: dict[str, dict[str, Any]] = {}

    def build(self, occurrences: list[dict[str, Any]]) -> tuple[DependencyGraph, dict[str, dict], list[dict]]:
        pattern_breaks = self._pattern_breaks(occurrences)
        for occurrence in sorted(occurrences, key=lambda item: (item["sheet_id"], item["cell_address"])):
            source = self.catalog.cell_node(
                occurrence["sheet_id"], occurrence["cell_address"], formula=occurrence["raw_formula"],
                pattern_id=occurrence["pattern_id"], normalized_hash=occurrence["normalized_hash"],
            )
            self.graph.add_node(source)
            try:
                ast = self.parser.parse(occurrence["raw_formula"])
                if occurrence["pattern_id"] not in self.ast_by_pattern:
                    nodes = list(ast.walk())
                    self.ast_by_pattern[occurrence["pattern_id"]] = {
                        "ast": ast.to_dict(), "ast_depth": ast.depth,
                        "normalized_hash": occurrence["normalized_hash"],
                        "function_count": sum(node.kind == "function" for node in nodes),
                        "reference_count": sum(node.kind == "reference" for node in nodes),
                        "range_count": sum(node.kind == "reference" and ":" in str(node.value) for node in nodes),
                        "dynamic_reference_count": sum(node.kind == "function" and str(node.value).upper() in {"INDIRECT", "OFFSET"} for node in nodes),
                    }
                self._resolve_ast(source, ast)
            except (FormulaParseError, ValueError, TypeError) as exc:
                self.graph.warnings.append({
                    "code": "FORMULA_PARSE_FAILED", "sheet_id": occurrence["sheet_id"],
                    "cell_address": occurrence["cell_address"], "message": str(exc),
                })
        return self.graph, self.ast_by_pattern, pattern_breaks

    def _resolve_ast(self, source: DependencyNode, ast) -> None:
        dynamic_functions = {
            str(node.value).upper() for node in ast.walk()
            if node.kind == "function" and str(node.value).upper() in {"INDIRECT", "OFFSET"}
        }
        for function in sorted(dynamic_functions):
            node_id = f"DYNAMIC:{source.node_id}:{function}"
            self.graph.add_node(DependencyNode(node_id, "RANGE", f"{function} dynamic target", source.sheet_id,
                                               metadata={"function": function, "source_formula": source.metadata.get("formula")}))
            self.graph.add_edge(DependencyEdge(source.node_id, node_id, category="DYNAMIC",
                                                resolution_status="AMBIGUOUS",
                                                metadata_json=json.dumps({"function": function}, sort_keys=True)))
        for node in ast.walk():
            if node.kind != "reference":
                continue
            self._resolve_reference(source, parse_reference(str(node.value)))
        for node in ast.walk():
            if node.kind == "error" and "#REF!" in str(node.value).upper():
                self._broken_edge(source, str(node.value))

    def _resolve_reference(self, source: DependencyNode, reference: FormulaReference) -> None:
        if reference.kind == "BROKEN":
            self._broken_edge(source, reference.raw)
            return
        if reference.workbook_name:
            self._external_edge(source, reference)
            return
        target_sheet = self.catalog.sheet_id(reference.sheet_name, source.sheet_id or "")
        if reference.sheet_name and not target_sheet:
            node_id = f"UNRESOLVED:SHEET:{_safe(reference.sheet_name)}"
            self.graph.add_node(DependencyNode(node_id, "SHEET", reference.sheet_name,
                                               metadata={"missing": True}))
            self.graph.add_edge(DependencyEdge(source.node_id, node_id, category="CROSS_SHEET",
                                                resolution_status="BROKEN",
                                                metadata_json=json.dumps({"reference": reference.raw}, sort_keys=True)))
            return
        if reference.kind == "CELL":
            target = self.catalog.cell_node(target_sheet or source.sheet_id or "", reference.start or "A1")
            self.graph.add_node(target)
            category = "CROSS_SHEET" if target.sheet_id != source.sheet_id else "DIRECT_CELL"
            self._edge(source, target, category, "RESOLVED", 1, reference)
        elif reference.kind in {"RANGE", "COLUMN_RANGE", "ROW_RANGE"}:
            self._range_edge(source, target_sheet or source.sheet_id or "", reference)
        elif reference.kind == "NAMED_RANGE":
            self._named_range_edge(source, reference)
        elif reference.kind == "TABLE_COLUMN":
            self._table_column_edge(source, reference)
        else:
            self._unresolved_edge(source, reference)

    def _edge(self, source: DependencyNode, target: DependencyNode, category: str,
              status: str, cardinality: int, reference: FormulaReference,
              dependency_type: str = "DEPENDS_ON") -> None:
        self.graph.add_edge(DependencyEdge(
            source.node_id, target.node_id, dependency_type, category, status, cardinality,
            json.dumps({"reference": reference.raw, "absolute_column": reference.absolute_column,
                        "absolute_row": reference.absolute_row}, sort_keys=True),
        ))

    def _range_edge(self, source: DependencyNode, sheet_id: str, reference: FormulaReference) -> None:
        max_row, max_column = self.catalog.sheet_dimensions.get(sheet_id, (1, 1))
        if reference.kind == "RANGE":
            min_col, min_row, max_col, end_row = range_boundaries(
                f"{_cell_address(reference.start or 'A1')}:{_cell_address(reference.end or 'A1')}"
            )
            max_row = end_row
        elif reference.kind == "COLUMN_RANGE":
            min_col = column_index_from_string((reference.start or "A").replace("$", ""))
            max_col = column_index_from_string((reference.end or "A").replace("$", ""))
            min_row = 1
        else:
            min_row = int((reference.start or "1").replace("$", ""))
            max_row = int((reference.end or "1").replace("$", ""))
            min_col, max_col = 1, max_column
        cardinality = max(1, (max_col - min_col + 1) * (max_row - min_row + 1))
        if cardinality <= settings.dependency_range_expansion_limit and reference.kind == "RANGE":
            for row in range(min_row, max_row + 1):
                for column in range(min_col, max_col + 1):
                    target = self.catalog.cell_node(sheet_id, f"{get_column_letter(column)}{row}")
                    self.graph.add_node(target)
                    category = "CROSS_SHEET" if sheet_id != source.sheet_id else "RANGE"
                    self._edge(source, target, category, "RESOLVED", 1, reference)
            return
        range_id = f"RANGE:{self.catalog.repository_id}:{sheet_id}:{_safe(reference.start or '')}:{_safe(reference.end or '')}"
        display = f"{self.catalog.sheet_names.get(sheet_id, sheet_id)}!{reference.start}:{reference.end}"
        target = self.graph.add_node(DependencyNode(
            range_id, "RANGE", display, sheet_id,
            metadata={"range_kind": reference.kind, "start": reference.start, "end": reference.end,
                      "logical_cardinality": cardinality},
        ))
        category = "CROSS_SHEET" if sheet_id != source.sheet_id else "RANGE"
        self._edge(source, target, category, "RESOLVED", cardinality, reference)

    def _named_range_edge(self, source: DependencyNode, reference: FormulaReference) -> None:
        name = (reference.name or reference.raw).casefold()
        definition = self.catalog.named_ranges.get(name)
        node_id = f"NAMED_RANGE:{self.catalog.repository_id}:{_safe(reference.name or reference.raw)}"
        named = self.graph.add_node(DependencyNode(node_id, "NAMED_RANGE", reference.name or reference.raw,
                                                   metadata={"definition": definition.get("range") if definition else None}))
        self._edge(source, named, "NAMED_RANGE", "RESOLVED" if definition else "BROKEN", 1, reference)
        if not definition:
            return
        resolved = parse_reference(str(definition.get("range") or "#REF!"))
        # Preserve the named semantic node while resolving its physical target.
        source_proxy = DependencyNode(named.node_id, named.node_type, named.display_name, named.sheet_id,
                                      metadata=named.metadata)
        self._resolve_reference(source_proxy, resolved)
        for edge in list(self.graph.edges.values()):
            if edge.source_node_id == named.node_id and edge.dependency_type == "DEPENDS_ON":
                self.graph.edges.pop(edge.key)
                self.graph.add_edge(DependencyEdge(edge.source_node_id, edge.target_node_id, "RESOLVES_TO",
                                                   edge.category, edge.resolution_status,
                                                   edge.logical_cardinality, edge.metadata_json))

    def _table_column_edge(self, source: DependencyNode, reference: FormulaReference) -> None:
        table_name = reference.table_name
        if not table_name:
            source_row, source_col = coordinate_to_tuple(source.cell_address or "A1")
            for (sheet_id, candidate_name), table in self.catalog.tables.items():
                if sheet_id != source.sheet_id:
                    continue
                try:
                    min_col, min_row, max_col, max_row = range_boundaries(table.get("range", "A1:A1"))
                    if min_row <= source_row <= max_row and min_col <= source_col <= max_col:
                        table_name = candidate_name
                        break
                except ValueError:
                    continue
        key = (source.sheet_id or "", (table_name or "").casefold())
        table = self.catalog.tables.get(key)
        node_id = f"TABLE_COLUMN:{self.catalog.repository_id}:{source.sheet_id}:{_safe(table_name or 'CURRENT')}:{_safe(reference.table_column or '')}"
        target = self.graph.add_node(DependencyNode(node_id, "TABLE_COLUMN",
                                                   f"{table_name or 'Current table'}.{reference.table_column}",
                                                   source.sheet_id, metadata={"table": table_name, "column": reference.table_column}))
        self._edge(source, target, "TABLE_COLUMN", "RESOLVED" if table else "AMBIGUOUS", 1, reference)

    def _external_edge(self, source: DependencyNode, reference: FormulaReference) -> None:
        filename = reference.workbook_name or "external"
        candidates = self.catalog.external_assets.get(filename.casefold(), [])
        status = "RESOLVED" if len(candidates) == 1 else "AMBIGUOUS" if len(candidates) > 1 else "EXTERNAL_UNAVAILABLE"
        node_id = f"EXTERNAL_WORKBOOK:{_safe(filename)}"
        target = self.graph.add_node(DependencyNode(node_id, "EXTERNAL_WORKBOOK", filename,
                                                   metadata={"candidates": candidates, "sheet": reference.sheet_name,
                                                             "address": reference.start or reference.name}))
        self._edge(source, target, "CROSS_WORKBOOK", status, 1, reference, "EXTERNAL_DEPENDENCY")

    def _broken_edge(self, source: DependencyNode, raw: str) -> None:
        node_id = f"BROKEN:{hashlib.sha256((source.node_id + raw).encode()).hexdigest()[:20].upper()}"
        target = self.graph.add_node(DependencyNode(node_id, "CELL", "Broken reference", source.sheet_id,
                                                   metadata={"reference": raw}))
        self.graph.add_edge(DependencyEdge(source.node_id, target.node_id, category="DIRECT_CELL",
                                            resolution_status="BROKEN",
                                            metadata_json=json.dumps({"reference": raw}, sort_keys=True)))

    def _unresolved_edge(self, source: DependencyNode, reference: FormulaReference) -> None:
        node_id = f"UNRESOLVED:{hashlib.sha256(reference.raw.encode()).hexdigest()[:20].upper()}"
        target = self.graph.add_node(DependencyNode(node_id, "RANGE", reference.raw,
                                                   metadata={"reference": reference.raw}))
        self._edge(source, target, "DYNAMIC", "UNRESOLVED", 1, reference)

    @staticmethod
    def _pattern_breaks(occurrences: list[dict[str, Any]]) -> list[dict[str, Any]]:
        groups = defaultdict(list)
        for item in occurrences:
            row, column = coordinate_to_tuple(item["cell_address"])
            groups[(item["sheet_id"], column)].append((row, item))
        breaks = []
        for (sheet_id, column), values in groups.items():
            values.sort()
            for index in range(1, len(values) - 1):
                previous, current, following = values[index - 1][1], values[index][1], values[index + 1][1]
                if previous["normalized_hash"] == following["normalized_hash"] != current["normalized_hash"]:
                    breaks.append({"sheet_id": sheet_id, "cell_address": current["cell_address"],
                                   "expected_pattern_hash": previous["normalized_hash"],
                                   "actual_pattern_hash": current["normalized_hash"]})
        return breaks
