"""Composable workbook and OpenXML inventory analyzers."""

import hashlib
import json
import re
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl.utils import get_column_letter

from .models import AnalysisContext, AnalyzerResult, AnalysisWarning

CELL_REF = re.compile(r"(?<![A-Z0-9_])(?P<sheet>(?:'[^']+'|[A-Za-z_][\w .]*)!)?(?P<col>\$?[A-Z]{1,3})(?P<row>\$?\d+)")
FUNCTION = re.compile(r"(?i)(?<![A-Z0-9_.])([A-Z][A-Z0-9_.]*)\s*\(")
EXTERNAL = re.compile(r"\[([^\]]+\.(?:xlsx|xlsm|xlsb|xls))\]", re.I)
VOLATILE = {"NOW", "TODAY", "RAND", "RANDBETWEEN", "OFFSET", "INDIRECT", "CELL", "INFO"}


def normalize_formula(formula: str) -> str:
    def replace(match):
        row = match.group("row")
        normalized_row = row if row.startswith("$") else "#"
        return f"{match.group('sheet') or ''}{match.group('col')}{normalized_row}"
    return CELL_REF.sub(replace, formula.upper().replace(" ", ""))


def compact_ranges(numbers: list[int]) -> list[str]:
    if not numbers:
        return []
    ranges = []
    start = previous = numbers[0]
    for number in numbers[1:]:
        if number != previous + 1:
            ranges.append(str(start) if start == previous else f"{start}:{previous}")
            start = number
        previous = number
    ranges.append(str(start) if start == previous else f"{start}:{previous}")
    return ranges


class WorkbookAnalyzer:
    name = "workbook-analyzer@1.0"

    def analyze(self, context: AnalysisContext) -> AnalyzerResult:
        result = AnalyzerResult(self.name)
        sheets = []
        formulas = defaultdict(list)
        object_records = []
        totals = Counter()
        for position, sheet in enumerate(context.workbook.worksheets):
            sheet_id = context.stable_sheets.get(sheet.title) or f"SHEET_{hashlib.sha256(sheet.title.encode()).hexdigest()[:16].upper()}"
            used = formula_count = constant_count = styled_blank = comments = hyperlinks = 0
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        styled_blank += int(cell.has_style)
                        continue
                    used += 1
                    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                        formula_count += 1
                        formulas[sheet_id].append((cell.coordinate, str(cell.value)))
                    else:
                        constant_count += 1
                    comments += int(cell.comment is not None)
                    hyperlinks += int(cell.hyperlink is not None)
                    if cell.comment:
                        object_records.append({"sheet_id": sheet_id, "object_type": "COMMENT", "name": cell.comment.author, "range": cell.coordinate, "details": {"text_hash": hashlib.sha256(cell.comment.text.encode()).hexdigest()}})
                    if cell.hyperlink:
                        target = str(cell.hyperlink.target or cell.hyperlink.location or "")
                        object_records.append({"sheet_id": sheet_id, "object_type": "HYPERLINK", "name": target[:240], "range": cell.coordinate, "details": {"external": bool(cell.hyperlink.target)}})
            hidden_rows = sorted(index for index, dimension in sheet.row_dimensions.items() if dimension.hidden)
            hidden_columns = sorted(dimension.min for dimension in sheet.column_dimensions.values() if dimension.hidden and dimension.min)
            table_count = len(sheet.tables)
            chart_count = len(sheet._charts)
            validations = len(sheet.data_validations.dataValidation)
            totals.update(used_cells=used, formula_cells=formula_count, constant_cells=constant_count,
                          blank_styled_cells=styled_blank, tables=table_count, charts=chart_count,
                          validations=validations, comments=comments, hyperlinks=hyperlinks)
            if sheet.sheet_state == "veryHidden":
                result.warnings.append(AnalysisWarning("VERY_HIDDEN_SHEET", f"{sheet.title} is very hidden.", sheet_id=sheet_id))
            for table in sheet.tables.values():
                object_records.append({"sheet_id": sheet_id, "object_type": "TABLE", "name": table.name, "range": table.ref, "details": {"style": getattr(table.tableStyleInfo, "name", None), "columns": [column.name for column in table.tableColumns]}})
            for item in sheet.data_validations.dataValidation:
                object_records.append({"sheet_id": sheet_id, "object_type": "VALIDATION", "name": item.type, "range": str(item.sqref), "details": {"formula1": item.formula1, "formula2": item.formula2, "allow_blank": item.allowBlank}})
            for merged in sheet.merged_cells.ranges:
                object_records.append({"sheet_id": sheet_id, "object_type": "MERGED_RANGE", "range": str(merged), "details": {}})
            for index, chart in enumerate(sheet._charts, 1):
                object_records.append({"sheet_id": sheet_id, "object_type": "CHART", "name": f"Chart {index}", "details": {"chart_type": chart.__class__.__name__, "title_present": chart.title is not None}})
            for index, image in enumerate(sheet._images, 1):
                object_records.append({"sheet_id": sheet_id, "object_type": "IMAGE", "name": f"Image {index}", "details": {"width": image.width, "height": image.height, "format": image.format}})
            used_range = f"A1:{get_column_letter(max(sheet.max_column, 1))}{max(sheet.max_row, 1)}"
            structure = {"name": sheet.title, "visibility": sheet.sheet_state, "rows": sheet.max_row, "columns": sheet.max_column, "tables": sorted(sheet.tables)}
            sheets.append({"sheet_id": sheet_id, "sheet_name": sheet.title, "position": position,
                           "visibility": sheet.sheet_state.upper(), "max_row": sheet.max_row,
                           "max_column": sheet.max_column, "used_range": used_range,
                           "used_cells": used, "formula_cells": formula_count,
                           "constant_cells": constant_count, "blank_styled_cells": styled_blank,
                           "merged_ranges": len(sheet.merged_cells.ranges), "hidden_rows": len(hidden_rows),
                           "hidden_columns": len(hidden_columns), "hidden_row_ranges": compact_ranges(hidden_rows),
                           "hidden_column_ranges": compact_ranges(hidden_columns), "tables": table_count,
                           "charts": chart_count, "pivots": len(sheet._pivots), "validations": validations,
                           "comments": comments, "hyperlinks": hyperlinks,
                           "structure_hash": hashlib.sha256(json.dumps(structure, sort_keys=True).encode()).hexdigest()})
        result.metrics = dict(totals)
        named_range_count = 0
        for defined_name in context.workbook.defined_names.values():
            named_range_count += 1
            object_records.append({"object_type": "NAMED_RANGE", "name": defined_name.name,
                                   "range": defined_name.attr_text,
                                   "details": {"hidden": bool(defined_name.hidden), "local_sheet_id": defined_name.localSheetId}})
        result.metrics.update(sheet_count=len(sheets), visible_sheets=sum(s["visibility"] == "VISIBLE" for s in sheets),
                              hidden_sheets=sum(s["visibility"] == "HIDDEN" for s in sheets),
                              very_hidden_sheets=sum(s["visibility"] == "VERYHIDDEN" for s in sheets),
                              named_ranges=named_range_count)
        result.records = {"sheets": sheets, "formula_cells": [{"sheet_id": key, "items": value} for key, value in formulas.items()], "objects": object_records}
        return result


class FormulaAnalyzer:
    name = "formula-analyzer@1.0"

    def analyze(self, context: AnalysisContext, workbook_result: AnalyzerResult) -> AnalyzerResult:
        result = AnalyzerResult(self.name)
        patterns = {}
        occurrences = []
        function_counts = Counter()
        for sheet in workbook_result.records["formula_cells"]:
            for address, raw in sheet["items"]:
                normalized = normalize_formula(raw)
                digest = hashlib.sha256(normalized.encode()).hexdigest()
                functions = sorted(set(match.upper() for match in FUNCTION.findall(raw)))
                function_counts.update(FUNCTION.findall(raw.upper()))
                external = bool(EXTERNAL.search(raw))
                cross_sheet = "!" in raw
                volatile = bool(set(functions) & VOLATILE)
                pattern = patterns.setdefault(digest, {"pattern_id": f"FPT_{digest[:20].upper()}", "normalized_formula": normalized,
                    "normalized_hash": digest, "occurrence_count": 0, "function_count": len(functions), "functions": functions,
                    "cross_sheet": cross_sheet, "external": external, "volatile": volatile})
                pattern["occurrence_count"] += 1
                occurrences.append({"pattern_id": pattern["pattern_id"], "sheet_id": sheet["sheet_id"], "cell_address": address, "raw_formula": raw})
        result.metrics = {"total_formulas": len(occurrences), "unique_patterns": len(patterns),
                          "cross_sheet_formulas": sum("!" in item["raw_formula"] for item in occurrences),
                          "external_formulas": sum(bool(EXTERNAL.search(item["raw_formula"])) for item in occurrences),
                          "volatile_formulas": sum(any(fn in item["raw_formula"].upper() for fn in VOLATILE) for item in occurrences),
                          "top_functions": function_counts.most_common(20)}
        result.records = {"patterns": list(patterns.values()), "occurrences": occurrences}
        return result


class OpenXMLAnalyzer:
    name = "openxml-analyzer@1.0"

    def analyze(self, context: AnalysisContext) -> AnalyzerResult:
        result = AnalyzerResult(self.name)
        if context.file_type == "csv":
            return result
        objects, links, connections = [], [], []
        with zipfile.ZipFile(context.file_path) as package:
            names = set(package.namelist())
            for name in sorted(names):
                if name.startswith("xl/externalLinks/externalLink") and name.endswith(".xml"):
                    root = ET.fromstring(package.read(name))
                    workbook_name = next((node.text for node in root.iter() if node.tag.endswith("sheetName")), None)
                    links.append({"source_euc_reference": workbook_name or name, "relationship_type": "EXTERNAL_WORKBOOK"})
                if name.startswith("xl/pivotTables/pivotTable") and name.endswith(".xml"):
                    root = ET.fromstring(package.read(name))
                    objects.append({"object_type": "PIVOT_TABLE", "name": root.attrib.get("name"), "details": {"part": name}})
            if "xl/connections.xml" in names:
                root = ET.fromstring(package.read("xl/connections.xml"))
                for node in root.iter():
                    if node.tag.endswith("connection"):
                        raw = json.dumps(node.attrib)
                        credential_present = bool(re.search(r"password|pwd|credential", raw, re.I))
                        connections.append({"name": node.attrib.get("name"), "type": node.attrib.get("type"),
                                            "credential_present": credential_present,
                                            "details": {key: "[REDACTED]" if re.search(r"password|pwd|credential", key, re.I) else value for key, value in node.attrib.items()}})
            power_query_parts = [name for name in names if "customXml" in name or "connections" in name and name != "xl/connections.xml"]
            if power_query_parts:
                objects.append({"object_type": "POWER_QUERY", "name": "Power Query package metadata", "details": {"part_count": len(power_query_parts)}})
            if "xl/vbaProject.bin" in names:
                payload = package.read("xl/vbaProject.bin")
                text_fragments = re.findall(rb"[A-Za-z_][A-Za-z0-9_]{3,80}", payload)
                tokens = {token.decode("latin-1", errors="ignore") for token in text_fragments}
                auto_events = sorted(name for name in tokens if name.lower() in {
                    "auto_open", "workbook_open", "workbook_beforeclose", "worksheet_change"
                })
                module_names = sorted(name for name in tokens if name.lower().startswith(("module", "sheet", "thisworkbook")))[:100]
                objects.append({"object_type": "VBA_PROJECT", "name": "vbaProject.bin", "details": {
                    "size_bytes": len(payload), "project_hash": hashlib.sha256(payload).hexdigest(),
                    "module_names": module_names, "module_count": len(module_names),
                    "auto_exec_events": auto_events, "execution": "DISABLED"}})
                result.warnings.append(AnalysisWarning("VBA_PRESENT", "VBA content was detected and statically inventoried; it was not executed."))
        result.metrics = {"external_links": len(links), "connections": len(connections),
                          "pivots": sum(item["object_type"] == "PIVOT_TABLE" for item in objects),
                          "power_queries": sum(item["object_type"] == "POWER_QUERY" for item in objects),
                          "vba_present": any(item["object_type"] == "VBA_PROJECT" for item in objects)}
        result.records = {"objects": objects, "external_links": links, "connections": connections}
        return result
