"""Branch workbook generation kept separate from HTTP route handling."""

import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..config import settings
from ..database import _get_connection, create_working_copy
from ..excel.identity import semantic_snapshot
from ..openxml_injector import inject_taskpane_manifest


def _write_branch_xlsx(table_id: str, destination: Path, include_data: bool = True) -> None:
    conn = _get_connection()
    try:
        snapshot = semantic_snapshot(conn, table_id)
    finally:
        conn.close()
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names: set[str] = set()
    for sheet_index, semantic_sheet in enumerate(snapshot.get("sheets", [])):
        base_name = str(semantic_sheet.get("name") or f"Sheet{sheet_index + 1}")[:31]
        sheet_name = base_name
        suffix = 2
        while sheet_name.casefold() in used_names:
            tail = f"_{suffix}"
            sheet_name = f"{base_name[:31 - len(tail)]}{tail}"
            suffix += 1
        used_names.add(sheet_name.casefold())
        sheet = workbook.create_sheet(sheet_name)
        columns = sorted(semantic_sheet.get("columns", []), key=lambda item: item["position"])
        for column_index, column in enumerate(columns, 1):
            cell = sheet.cell(row=1, column=column_index, value=column["name"])
            cell.fill = PatternFill("solid", fgColor="161B22")
            cell.font = Font(color="FFFFFF", bold=True)
        if include_data:
            rows = sorted(semantic_sheet.get("rows", []), key=lambda item: item["position"])
            for row_index, row in enumerate(rows, 2):
                for column_index, column in enumerate(columns, 1):
                    column_id = column["column_id"]
                    value = row.get("formulas", {}).get(column_id)
                    if value is None:
                        value = row.get("values", {}).get(column_id)
                    sheet.cell(row=row_index, column=column_index, value=value)
            if columns and rows:
                table = Table(
                    displayName=f"GitWalkData{sheet_index + 1}", ref=sheet.dimensions
                )
                table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                sheet.add_table(table)
        else:
            if columns:
                msg = "[LOCKED] Verify Path & Sign In via Git Walk Taskpane to load data"
                placeholder = sheet.cell(row=2, column=1, value=msg)
                placeholder.font = Font(italic=True, color="58A6FF")
        sheet.freeze_panes = "A2"
    if not workbook.worksheets:
        workbook.create_sheet("Sheet1")
    workbook.save(destination)


def issue_branch_workbook(
    main_table_id: str,
    user_id: str,
    user_email: str,
    source_workbook: str | None = None,
    branch_mode: str = "continue",
    branch_id: str | None = None,
    local_file_path: str | None = None,
    local_target_dir: str | None = None,
) -> dict:
    """Create/reuse a personal branch and return a signed branch workbook."""
    identity = create_working_copy(
        main_table_id, user_id, user_email, branch_mode=branch_mode,
        branch_id=branch_id,
    )
    if not local_file_path and local_target_dir:
        filename = f"gitwalk_{identity['branch_name'].replace('/', '_')}.xlsx"
        local_file_path = str(Path(local_target_dir) / filename)
    if local_file_path:
        identity["local_file_path"] = local_file_path
        identity["required_role"] = "editor"
    directory = Path(tempfile.mkdtemp(prefix="gitwalk_working_copy_"))
    raw_path = directory / "branch.xlsx"
    output_path = directory / f"gitwalk_{identity['branch_id']}.xlsx"
    if source_workbook:
        raw_path.write_bytes(Path(source_workbook).read_bytes())
    else:
        _write_branch_xlsx(identity["table_id"], raw_path, include_data=False)
    inject_taskpane_manifest(
        input_xlsx_path=str(raw_path),
        output_xlsx_path=str(output_path),
        manifest_url=f"{settings.office_addin_url}/taskpane.html",
        table_id=identity["table_id"],
        metadata=identity,
    )
    return {**identity, "path": str(output_path)}


def export_branch_workbook(table_id: str, branch_id: str, branch_name: str) -> dict:
    """Render a complete read-only branch export from persisted semantic state."""
    directory = Path(tempfile.mkdtemp(prefix="gitwalk_branch_export_"))
    path = directory / f"{branch_name.replace('/', '_')}.xlsx"
    _write_branch_xlsx(table_id, path)
    return {"path": str(path), "branch_id": branch_id, "branch_name": branch_name}
